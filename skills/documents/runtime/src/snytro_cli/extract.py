# SPDX-License-Identifier: Apache-2.0
# Copyright 2026 Snytro — https://snytro.com
from __future__ import annotations

import csv
import hashlib
import io
import json
import re
from collections import Counter
from pathlib import Path


SUPPORTED = {".pdf", ".docx", ".txt", ".text", ".md", ".rst", ".csv", ".tsv", ".json", ".html", ".htm", ".xlsx"}

# Standard lexical stopwords prevent common function words dominating rank fusion.
# This is tokenization, not task/planner routing; names, numbers and negation remain.
SEARCH_STOPWORDS = frozenset("a an and are as at be been being by can could do does for from had has have how i in into is it its me much my of on or our should than that the their them these they this those to was we were what when where which who why will with would you your".split())


def discover(paths: list[str], recursive: bool = True) -> tuple[list[Path], list[dict]]:
    files, errors, seen = [], [], set()
    for raw in paths:
        try:
            root = Path(raw).expanduser().resolve(strict=True)
            if not root.is_file() and not root.is_dir():
                raise ValueError("Path must be a regular file or directory.")
            candidates = root.rglob("*") if root.is_dir() and recursive else root.glob("*") if root.is_dir() else [root]
            for item in candidates:
                # Named symlinks resolve above; discovered symlinks never expand folder scope.
                if item.is_symlink() or not item.is_file():
                    continue
                if item.suffix.lower() not in SUPPORTED:
                    if root.is_file():
                        errors.append({"path": str(item), "error": "Unsupported file type."})
                    continue
                resolved = item.resolve(strict=True)
                if root.is_dir() and not resolved.is_relative_to(root):
                    continue
                if resolved not in seen:
                    files.append(resolved)
                    seen.add(resolved)
        except (OSError, ValueError) as exc:
            errors.append({"path": raw, "error": str(exc)})
    return sorted(files), errors


def checksum(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def extract_units(path: Path) -> list[dict]:
    if path.stat().st_size > 128 * 1024 * 1024:
        raise ValueError("File exceeds the 128 MiB per-file extraction limit.")
    extension = path.suffix.lower()
    contents = []
    if extension == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(str(path))
        for number, page in enumerate(reader.pages, 1):
            contents.append((page.extract_text() or "", {"kind": "page", "page": number}))
    elif extension == ".docx":
        from docx import Document
        from docx.table import Table
        document = Document(path)
        # Preserve body order; Word layout does not expose reliable physical page numbers.
        text = []
        for element in document.iter_inner_content():
            if isinstance(element, Table):
                text.extend("\t".join(cell.text for cell in row.cells) for row in element.rows)
            else:
                text.append(element.text)
        contents.append(("\n".join(text), {"kind": "document"}))
    elif extension == ".xlsx":
        from openpyxl import load_workbook
        workbook = load_workbook(path, read_only=True, data_only=False)
        try:
            for sheet in workbook:
                rows = []
                start = 1
                for number, row in enumerate(sheet.iter_rows(values_only=True), 1):
                    rows.append("\t".join("" if value is None else str(value) for value in row))
                    if len(rows) == 100:
                        contents.append(("\n".join(rows), {"kind": "sheet", "sheet": sheet.title, "row_start": start, "row_end": number}))
                        rows, start = [], number + 1
                if rows:
                    contents.append(("\n".join(rows), {"kind": "sheet", "sheet": sheet.title, "row_start": start, "row_end": start + len(rows) - 1}))
        finally:
            workbook.close()
    else:
        raw = path.read_text(encoding="utf-8-sig")
        if extension in {".html", ".htm"}:
            from bs4 import BeautifulSoup
            soup = BeautifulSoup(raw, "html.parser")
            for node in soup(["script", "style"]):
                node.decompose()
            raw = soup.get_text("\n")
        elif extension == ".json":
            raw = json.dumps(json.loads(raw), ensure_ascii=False, indent=2)
        elif extension in {".csv", ".tsv"}:
            reader = csv.reader(io.StringIO(raw), delimiter="\t" if extension == ".tsv" else ",")
            raw = "\n".join("\t".join(row) for row in reader)
        contents.append((raw, {"kind": "document"}))
    units = [{"number": i, "content": text.replace("\x00", ""), "locator": locator} for i, (text, locator) in enumerate(contents, 1)]
    if not any(unit["content"].strip() for unit in units):
        raise ValueError("No extractable text. Scanned PDFs require OCR; OCR is not included in this release.")
    return units


def terms(text: str) -> list[str]:
    return re.findall(r"[^\W_]+", text.lower(), flags=re.UNICODE)


def overview(text: str) -> dict:
    # Bootstrap excerpts only; the skill replaces them with host-generated analysis.
    sentences = [part.strip() for part in re.split(r"(?<=[.!?])\s+|\n+", text) if part.strip()]
    return {
        "tagline": (sentences[0] if sentences else "")[:180],
        "summary": "\n".join(sentences[:8])[:2400],
        "keywords": [word for word, _ in Counter(word for word in terms(text) if len(word) > 3).most_common(20)],
    }


def text_chunks(text: str) -> list[str]:
    return [text[start:start + 1600] for start in range(0, len(text), 1400) if text[start:start + 1600].strip()]
